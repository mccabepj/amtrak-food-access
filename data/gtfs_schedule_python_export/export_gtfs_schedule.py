#!/usr/bin/env python3
"""Export GTFS route/station arrival and dwell records to XLSX and CSV.
Requires Python 3.10+ and openpyxl. No web requests are made.
"""
import argparse
import csv
import io
import re
import zipfile
from collections import defaultdict
from datetime import date, datetime, time, timedelta, timezone
from pathlib import Path
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

HEADERS = [
    "Route", "Direction", "Station", "Arrival Time", "Dwell Minutes",
    "Departure Time", "Service / Train", "Stop Status", "Dwell Status",
    "Service Date", "Arrival Date", "Departure Date",
    "Arrival Day Offset", "Departure Day Offset", "Agency Timezone",
    "Effective Stop Timezone Metadata", "Time Precision", "Schedule Record Type",
    "Route ID", "Trip ID", "Stop ID", "Station ID", "Service ID", "Stop Sequence",
    "Route Type", "GTFS Arrival Time", "GTFS Departure Time",
    "Pickup Type", "Drop-off Type", "Direction Stop ID", "Direction Station ID"
]


def read_table(archive, filename, required=False, fields=()):
    matches = [n for n in archive.namelist()
               if not n.endswith("/") and not n.startswith("__MACOSX/")
               and n.rsplit("/", 1)[-1].lower() == filename]
    if len(matches) > 1:
        raise ValueError(f"Multiple {filename} files found. Use one GTFS feed.")
    if not matches:
        if required:
            raise ValueError(f"Missing {filename}.")
        return []
    with archive.open(matches[0]) as source:
        text = io.TextIOWrapper(source, encoding="utf-8-sig", newline="")
        reader = csv.DictReader(text, strict=True)
        if reader.fieldnames is None:
            if required:
                raise ValueError(f"{filename} is empty.")
            return []
        reader.fieldnames = [name.strip() for name in reader.fieldnames]
        missing = set(fields) - set(reader.fieldnames)
        if missing:
            raise ValueError(f"{filename} missing fields: {', '.join(sorted(missing))}")
        rows = []
        for row in reader:
            if None in row or any(value is None for value in row.values()):
                raise ValueError(f"Malformed row in {filename}, line {reader.line_num}.")
            if any(value.strip() for value in row.values()):
                rows.append(row)
        return rows


def unique_index(rows, key, filename):
    result = {}
    for row in rows:
        value = row.get(key, "")
        if not value or value in result:
            raise ValueError(f"Blank or duplicate {key} in {filename}: {value!r}")
        result[value] = row
    return result


def seconds(value):
    match = re.fullmatch(r"(\d+):(\d{2}):(\d{2})", value or "")
    if not match:
        return None
    hours, minutes, secs = map(int, match.groups())
    if minutes >= 60 or secs >= 60:
        return None
    return hours * 3600 + minutes * 60 + secs


def clock(value):
    total = seconds(value)
    if total is None:
        return "", None
    hours, remainder = divmod(total % 86400, 3600)
    minutes, secs = divmod(remainder, 60)
    return f"{hours:02}:{minutes:02}:{secs:02}", total // 86400


def rail(route):
    try:
        code = int(route.get("route_type", ""))
        return code == 2 or 100 <= code <= 117
    except ValueError:
        return False


def parent_stop(stop_id, stops):
    current = stops[stop_id]
    seen = {stop_id}
    while current.get("parent_station"):
        parent = current["parent_station"]
        if parent not in stops or parent in seen:
            break
        current = stops[parent]
        seen.add(parent)
    return current


def active_services(day, calendar, exceptions):
    key = day.strftime("%Y%m%d")
    weekday = ("monday", "tuesday", "wednesday", "thursday", "friday", "saturday", "sunday")[day.weekday()]
    result = {row["service_id"] for row in calendar
              if row["start_date"] <= key <= row["end_date"] and row.get(weekday) == "1"}
    for row in exceptions:
        if row["date"] == key:
            if row["exception_type"] == "1":
                result.add(row["service_id"])
            elif row["exception_type"] == "2":
                result.discard(row["service_id"])
    return result


def dated_time(day, value, zone):
    total = seconds(value)
    if total is None:
        return "", ""
    # GTFS origin of time: service-day local noon minus 12 elapsed hours.
    noon = datetime.combine(day, time(12), ZoneInfo(zone)).astimezone(timezone.utc)
    instant = noon - timedelta(hours=12) + timedelta(seconds=total)
    local = instant.astimezone(ZoneInfo(zone))
    return local.date().isoformat(), local.strftime("%H:%M:%S")


def export_rows(args):
    with zipfile.ZipFile(args.gtfs_zip) as archive:
        stops_rows = read_table(archive, "stops.txt", True, ("stop_id", "stop_name"))
        routes_rows = read_table(archive, "routes.txt", True, ("route_id", "route_type"))
        trips_rows = read_table(archive, "trips.txt", True, ("trip_id", "route_id", "service_id"))
        times_rows = read_table(archive, "stop_times.txt", True, ("trip_id", "stop_id", "stop_sequence"))
        agencies = read_table(archive, "agency.txt", True, ("agency_name", "agency_timezone"))
        calendar = read_table(archive, "calendar.txt", False, ("service_id", "start_date", "end_date"))
        exceptions = read_table(archive, "calendar_dates.txt", False, ("service_id", "date", "exception_type"))
        frequencies = read_table(archive, "frequencies.txt")
    stops = unique_index(stops_rows, "stop_id", "stops.txt")
    routes_all = unique_index(routes_rows, "route_id", "routes.txt")
    trips_all = unique_index(trips_rows, "trip_id", "trips.txt")
    routes = {key: row for key, row in routes_all.items() if args.include_non_rail or rail(row)}
    if not routes:
        raise ValueError("No matching routes found. Default is rail only.")
    active = None
    if args.date:
        if not calendar and not exceptions:
            raise ValueError("A date-specific export requires calendar.txt or calendar_dates.txt.")
        active = active_services(args.date, calendar, exceptions)
    trips = {key: trip for key, trip in trips_all.items()
             if trip["route_id"] in routes and (active is None or trip["service_id"] in active)}
    times = defaultdict(list)
    for row in times_rows:
        if row["trip_id"] not in trips:
            continue
        if row["stop_id"] not in stops:
            raise ValueError(f"Stop {row['stop_id']} is referenced but not in stops.txt.")
        try:
            sequence = int(row["stop_sequence"])
            if sequence < 0:
                raise ValueError()
        except ValueError:
            raise ValueError(f"Invalid stop_sequence in trip {row['trip_id']}.")
        row["_sequence"] = sequence
        times[row["trip_id"]].append(row)
    frequency_ids = {row.get("trip_id") for row in frequencies}
    agency_by_id = {row.get("agency_id", ""): row for row in agencies}
    sole_zone = next(iter({a["agency_timezone"] for a in agencies})) if len({a["agency_timezone"] for a in agencies}) == 1 else None
    notes = [
        ("Source ZIP", str(args.gtfs_zip.resolve())),
        ("Scope", "All route types" if args.include_non_rail else "Rail only (route_type 2 or 100–117)"),
        ("Export mode", f"Service date {args.date.isoformat()}" if args.date else "Trip-stop schedule templates; NOT one row per operating day"),
        ("Direction", "Final scheduled stop of each specific trip, using its parent-station name when available; not direction_id or headsign"),
        ("Station", "Parent-station name when available; Stop ID retains the actual scheduled stop/platform"),
        ("Times", "Agency timezone, not station-local time. Effective Stop Timezone Metadata is informational only."),
        ("Default time display", "24-hour clock with separate Arrival/Departure Day Offset. Raw extended-hour GTFS times are retained."),
        ("Dated display", "With --date, explicit dates and times are calculated using GTFS local-noon-minus-12h service-day timing; raw day offsets remain available."),
        ("Dwell Minutes", "Numeric departure minus arrival for intermediate stops. Endpoints/unknown/invalid times are blank; zero is recorded zero."),
        ("Frequency records", "Frequency-based templates are included and flagged, NOT expanded into individual departures. See Frequencies sheet."),
        ("Schedule validity", "Default template rows are not daily departures or frequency counts. See Service ID, Calendar, and Calendar exceptions."),
        ("Stop coverage", "Only scheduled trip-stop calls; unused stops without retained trips are not exported."),
        ("CSV safety", "Strings beginning with formula characters are prefixed with an apostrophe in CSV; XLSX stores strings as literal text."),
        ("Endpoint interpretation", "Blank dwell at trip origin/destination does not describe a physical train's terminal layover."),
    ]
    records = []
    for trip_id, trip in trips.items():
        calls = sorted(times.get(trip_id, []), key=lambda row: row["_sequence"])
        if not calls:
            continue
        if len({r["_sequence"] for r in calls}) != len(calls):
            raise ValueError(f"Duplicate stop_sequence in trip {trip_id}.")
        final_stop = stops[calls[-1]["stop_id"]]
        final_station = parent_stop(final_stop["stop_id"], stops)
        route = routes[trip["route_id"]]
        agency = agency_by_id.get(route.get("agency_id", ""))
        zone = agency.get("agency_timezone") if agency else sole_zone
        if not zone:
            raise ValueError(f"Cannot resolve agency timezone for route {trip['route_id']}.")
        if args.date:
            try:
                ZoneInfo(zone)
            except ZoneInfoNotFoundError:
                raise ValueError(f"Timezone {zone!r} unavailable. Install tzdata: py -m pip install tzdata")
        for i, call in enumerate(calls):
            stop = stops[call["stop_id"]]
            station = parent_stop(stop["stop_id"], stops)
            arrival, arrival_offset = clock(call.get("arrival_time", ""))
            departure, departure_offset = clock(call.get("departure_time", ""))
            a, d = seconds(call.get("arrival_time", "")), seconds(call.get("departure_time", ""))
            endpoint = i == 0 or i == len(calls) - 1
            status = "Origin and destination" if len(calls) == 1 else "Origin" if i == 0 else "Destination" if i == len(calls) - 1 else "Intermediate"
            dwell = None
            if endpoint:
                dwell_status = "Trip endpoint — excluded"
            elif a is None or d is None:
                dwell_status = "Missing or invalid arrival/departure"
            elif d < a:
                dwell_status = "Invalid: departure before arrival"
            else:
                dwell = round((d - a) / 60, 6)
                dwell_status = "Zero recorded dwell" if dwell == 0 else "Scheduled dwell"
            arrival_date = departure_date = ""
            if args.date:
                arrival_date, arrival = dated_time(args.date, call.get("arrival_time", ""), zone)
                departure_date, departure = dated_time(args.date, call.get("departure_time", ""), zone)
            records.append([
                route.get("route_long_name") or route.get("route_short_name") or route["route_id"],
                final_station.get("stop_name") or final_stop["stop_name"], station.get("stop_name") or stop["stop_name"],
                arrival, dwell, departure, trip.get("trip_short_name") or trip_id, status, dwell_status,
                args.date.isoformat() if args.date else "", arrival_date, departure_date,
                arrival_offset, departure_offset, zone, station.get("stop_timezone") or zone,
                "Estimated" if call.get("timepoint") == "0" else "Exact/default timepoint",
                "Frequency-based template; not expanded" if trip_id in frequency_ids else "Fixed-time trip template" if not args.date else "Scheduled call on service date",
                trip["route_id"], trip_id, stop["stop_id"], station["stop_id"], trip["service_id"], call["_sequence"],
                route["route_type"], call.get("arrival_time", ""), call.get("departure_time", ""),
                call.get("pickup_type") or "0", call.get("drop_off_type") or "0", final_stop["stop_id"], final_station["stop_id"]
            ])
    records.sort(key=lambda row: (row[0].casefold(), row[1].casefold(), row[19], row[23]))
    notes.extend([("Exported rows", len(records)), ("Routes represented", len({row[18] for row in records})),
                  ("Station IDs represented", len({row[21] for row in records}))])
    retained_services = {row[22] for row in records}
    tables = {
        "Calendar": [r for r in calendar if r["service_id"] in retained_services],
        "Calendar exceptions": [r for r in exceptions if r["service_id"] in retained_services],
        "Frequencies": [r for r in frequencies if r.get("trip_id") in trips],
    }
    return records, notes, tables


def write_workbook(path, records, notes, tables):
    try:
        from openpyxl import Workbook
        from openpyxl.cell import WriteOnlyCell
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise ValueError("Missing openpyxl. Install: py -m pip install openpyxl tzdata")
    workbook = Workbook(write_only=True)
    def sheet(title, headers, rows):
        ws = workbook.create_sheet(title)
        ws.freeze_panes = "A2"
        for i, name in enumerate(headers, 1):
            ws.column_dimensions[get_column_letter(i)].width = min(48, max(17, len(name) + 3))
        cells = []
        for name in headers:
            cell = WriteOnlyCell(ws, name)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="315F52")
            cells.append(cell)
        ws.append(cells)
        count = 1
        for row in rows:
            cells = []
            for value in row:
                cell = WriteOnlyCell(ws, value)
                if isinstance(value, str):
                    cell.data_type = "s"
                if isinstance(value, float):
                    cell.number_format = "0.##"
                cells.append(cell)
            ws.append(cells)
            count += 1
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{count}"
    page_size = 1048575
    for offset in range(0, max(len(records), 1), page_size):
        number = offset // page_size + 1
        sheet("Results" if number == 1 else f"Results {number}", HEADERS, records[offset:offset + page_size])
    sheet("Notes", ["Item", "Explanation"], notes)
    for title, rows in tables.items():
        if rows:
            headers = list(rows[0])
            for offset in range(0, len(rows), page_size):
                page = offset // page_size + 1
                sheet(title if page == 1 else f"{title} {page}", headers,
                      ([row.get(h, "") for h in headers] for row in rows[offset:offset + page_size]))
    workbook.save(path)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("gtfs_zip", type=Path, help="Path to the GTFS ZIP; leave it zipped")
    parser.add_argument("--output", type=Path, default=Path("exports"), help="Output directory (default: exports)")
    parser.add_argument("--date", type=date.fromisoformat, help="Optional SERVICE date YYYY-MM-DD; otherwise export templates")
    parser.add_argument("--include-non-rail", action="store_true", help="Include buses and other modes; default is rail only")
    parser.add_argument("--overwrite", action="store_true", help="Allow replacement of existing export files")
    args = parser.parse_args()
    try:
        if not args.gtfs_zip.is_file():
            raise ValueError(f"GTFS ZIP not found: {args.gtfs_zip}")
        suffix = args.date.isoformat() if args.date else "templates"
        base = f"route_station_schedule_{suffix}"
        xlsx_path, csv_path = args.output / f"{base}.xlsx", args.output / f"{base}.csv"
        if not args.overwrite and (xlsx_path.exists() or csv_path.exists()):
            raise ValueError("Export files already exist. Use another output folder or add --overwrite.")
        records, notes, tables = export_rows(args)
        args.output.mkdir(parents=True, exist_ok=True)
        write_workbook(xlsx_path, records, notes, tables)
        with csv_path.open("w", encoding="utf-8-sig", newline="") as output:
            writer = csv.writer(output)
            writer.writerow(HEADERS)
            for row in records:
                writer.writerow(["'" + value if isinstance(value, str) and value.startswith(("=", "+", "-", "@", "\t", "\r")) else value for value in row])
        print(f"Exported {len(records):,} scheduled-call records.")
        print(f"Workbook: {xlsx_path.resolve()}")
        print(f"CSV:      {csv_path.resolve()}")
        if not records:
            print("No matching calls. Check route types or the selected service date.")
        print("Times use agency timezone. Read the workbook Notes sheet before analysis.")
    except (ValueError, OSError, zipfile.BadZipFile, csv.Error) as error:
        parser.exit(1, f"Error: {error}\n")


if __name__ == "__main__":
    main()
